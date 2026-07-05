from pptx import Presentation

prs = Presentation('Presentacion_TFG_Jacob_Pino.pptx')
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text.strip())
